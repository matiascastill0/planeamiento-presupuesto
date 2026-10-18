import os as _os; RAIZ = _os.environ.get('PPTO27_RAIZ', _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))) + '/'  # raíz del proyecto
"""Hoja 'Leyenda Compañías': vehículos de la unidad, agrupados por país (maestro de Compañías y Sociedades)."""
import openpyxl, pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
ROOT = RAIZ
PAIS_TXT = {'CHILE': 'Chile', 'COLOMBIA': 'Colombia', 'PERU': 'Perú', 'PANAMA': 'Panamá', 'USA': 'USA', 'MEXICO': 'México'}
MON = {'CHILE': 'CLP', 'COLOMBIA': 'COP', 'PERU': 'PEN', 'PANAMA': 'USD', 'USA': 'USD', 'MEXICO': 'USD'}

def maestro():
    ws = openpyxl.load_workbook(ROOT + '02. Insumos/06. Maestros/Compañías y Sociedades de Credicorp.xlsx').active
    m = {}
    for cod, abr, veh, pais in ws.iter_rows(min_row=4, min_col=2, max_col=5, values_only=True):
        if abr and abr != 'Holding':
            m[abr] = (cod, veh)
    return m
MAESTRO = maestro()
def _vehiculos_pais():
    ws = openpyxl.load_workbook(ROOT + '02. Insumos/06. Maestros/Compañías y Sociedades de Credicorp.xlsx').active
    m = {'Colombia': 'COLOMBIA', 'Chile': 'CHILE', 'Perú': 'PERU', 'México': 'MEXICO', 'US': 'USA', 'Panamá': 'PANAMA'}
    out = {}
    for cod, abr, veh, pais in ws.iter_rows(min_row=4, min_col=2, max_col=5, values_only=True):
        if abr and abr != 'Holding' and pais in m:
            out.setdefault(m[pais], []).append(abr)
    return out
VEHICULOS_PAIS = _vehiculos_pais()
NOMBRES = {k: v[1] for k, v in MAESTRO.items()}

AZUL = '1F3864'; AZUL2 = 'D9E1F2'; GRIS = 'F2F2F2'; VERDE = 'E2EFDA'
fina = Side(style='thin', color='BFBFBF')
B = Border(left=fina, right=fina, top=fina, bottom=fina)
FN = 'Aptos Narrow'

def agregar_hoja(wb, unidad, pais_archivo, R_u, P_u):
    """R_u / P_u: BD real y ppto de la unidad en TODOS los países."""
    if 'Leyenda Compañías' in wb.sheetnames:
        del wb['Leyenda Compañías']
    ws = wb.create_sheet('Leyenda Compañías', index=1)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = AZUL
    for col, w in zip('ABCDEFGH', [2, 10, 13, 58, 20, 20, 16, 2]):
        ws.column_dimensions[col].width = w
    ws['B2'] = f'Leyenda de compañías · {unidad}'
    ws['B2'].font = Font(name=FN, size=16, bold=True, color=AZUL)
    ws['B3'] = ('Vehículos (sociedades) donde la unidad registra gasto en Real 2026 (Ene–Jul) o Ppto 2026, '
                'agrupados por país. Fuente: maestro "Compañías y Sociedades de Credicorp".')
    ws['B3'].font = Font(name=FN, size=10, italic=True, color='595959')
    ws['B4'] = f'Este archivo corresponde a: {PAIS_TXT[pais_archivo]}  (sección resaltada en verde)'
    ws['B4'].font = Font(name=FN, size=10, bold=True, color='375623')
    real = R_u.groupby(['Pais', 'Compañía']).Monto.sum()
    ppto = P_u.groupby(['Pais', 'Compañía']).Monto.sum()
    paises = sorted(set(real.index.get_level_values(0)) | set(ppto.index.get_level_values(0)),
                    key=lambda p: (p != pais_archivo, PAIS_TXT[p]))
    r = 6
    for p in paises:
        actual = p == pais_archivo
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        c = ws.cell(r, 2, f'{PAIS_TXT[p]}' + ('   ◀ este archivo' if actual else ''))
        c.font = Font(name=FN, size=12, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='375623' if actual else AZUL)
        c.alignment = Alignment(vertical='center', indent=1)
        ws.row_dimensions[r].height = 22
        r += 1
        for j, h in enumerate(['Código', 'Abreviatura', 'Vehículo (razón social)', f'Real 2026 Ene–Jul ({MON[p]})', f'Ppto 2026 ({MON[p]})', '¿Tiene gasto?'], start=2):
            c = ws.cell(r, j, h)
            c.font = Font(name=FN, size=10, bold=True, color=AZUL)
            c.fill = PatternFill('solid', fgColor=AZUL2); c.border = B
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.row_dimensions[r].height = 28
        r += 1
        comps = sorted(set(real.get(p, pd.Series(dtype=float)).index) | set(ppto.get(p, pd.Series(dtype=float)).index))
        for i, comp in enumerate(comps):
            cod, veh = MAESTRO.get(comp, (None, 'No figura en el maestro de compañías – por confirmar'))
            rv = real.get((p, comp)); pv = ppto.get((p, comp))
            tiene = ' y '.join(x for x, v in (('Real', rv), ('Ppto', pv)) if v is not None and abs(v) > 0.005) or '—'
            vals = [cod, comp, veh, rv, pv, tiene]
            fill = PatternFill('solid', fgColor=(VERDE if actual else GRIS) if i % 2 == 0 else 'FFFFFF')
            for j, v in enumerate(vals, start=2):
                c = ws.cell(r, j, v)
                c.font = Font(name=FN, size=10, bold=(j == 3), color=('C00000' if (j == 4 and comp not in MAESTRO) else '000000'))
                c.fill = fill; c.border = B
                c.alignment = Alignment(horizontal='left' if j == 4 else 'center', vertical='center')
                if j in (5, 6):
                    c.number_format = '#,##0;(#,##0);"-"'; c.alignment = Alignment(horizontal='right', vertical='center')
            r += 1
        r += 1
    ws.cell(r, 2, 'Notas').font = Font(name=FN, size=11, bold=True, color=AZUL)
    notas = ['Se excluyen las compañías INTERCOM, ELIM_SC y ELIM_CCSPP (no son registros de la contabilidad de las sociedades).',
             'En Panamá, "Commision Paid" y "TRASLADOS" se presentan dentro de BNK (ASB Bank Corp.).',
             'Montos en moneda local de cada país; negativo = gasto. La columna "Vehículo" de la hoja Ppto27 usa estas abreviaturas.']
    for n in notas:
        r += 1
        ws.cell(r, 2, '•  ' + n).font = Font(name=FN, size=10, color='404040')
    ws.freeze_panes = 'A6'
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1
    return ws
