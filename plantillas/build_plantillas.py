import os as _os; RAIZ = _os.environ.get('PPTO27_RAIZ', _os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))) + '/'  # raíz del proyecto
"""Genera las plantillas Ppto27 por unidad (UdN / Staff) y país a partir de la plantilla AM Chile."""
import pandas as pd, numpy as np, os, sys, copy, datetime as dt, re
import openpyxl, zipfile
from lxml import etree
from openpyxl.utils import get_column_letter
import pivot_rebuild
import leyenda
UNIDADES_LEYENDA = {'WM', 'JV CND', 'NNFF'}
from openpyxl.worksheet.table import Table
from agregar import preparar, PAIS_TXT, CUENTAS_CLIENTE
from compactar import agrupar_amplios
from prov_norm import exigir_sin_varios
from homologacion import Homologador

ROOT = RAIZ
TEMPLATE = ROOT + '02. Insumos/07. Plantillas ejemplo/Ppto27 - AM Chile - OFICIAL 1.xlsx'
OUT = ROOT + '99. Archivo/Versiones anteriores (salida)/'
HDR = 9                      # fila de encabezados de Tabla1
FIRST = HDR + 1
MES_COL = {m: 13 + m for m in range(1, 13)}   # Ene-26 = col N (14)

VEHICULOS = {
    'CHILE': {'CCCH': 'Credicorp Capital Chile S.A.', 'CCAF': 'Credicorp Capital Asesorías Financieras S.A',
              'CCAI': 'Credicorp Capital S.A. Administradora de Inversiones', 'CCCB': 'Credicorp Capital S.A. Corredores de Bolsa',
              'CAGF': 'Credicorp Capital Asset Magament S.A. Administradora General'},
    'COLOMBIA': {'SC': 'Credicorp Capital Colombia S.A.', 'FI': 'Fiduciaria Credicorp Capital S.A.',
                 'ELIM_SC': 'Eliminaciones SC', 'ELIM_CCSPP': 'Eliminaciones CCSPP'},
    'PERU': {'CCB': 'Credicorp Capital Sociedad Agente de Bolsa S.A.', 'CCF': 'Credicorp Capital S.A. Sociedad Administradora de Fondos',
             'CCP': 'Credicorp Capital Perú S.A.A.', 'CCSF': 'Credicorp Capital Servicios Financieros S.A.',
             'CCT': 'Credicorp Capital Sociedad Titulizadora S.A.', 'LTD': 'Credicorp Capital Ltd.'},
    'USA': {'LLC': 'Credicorp Capital LLC', 'ADV': 'Credicorp Capital Advisors LLC'},
    'PANAMA': {}, 'MEXICO': {},
}

def ceco_norm(c):
    s = str(c).strip()
    s = re.sub(r'\.0$', '', s)
    if re.fullmatch(r'211\d{3}', s):      # CeCo Colombia con prefijo de sociedad (211601 -> 601)
        s = s[3:]
    return s

def copiar_estilo(src, dst):
    dst._style = copy.copy(src._style)

def construir(unidad, tipo, pais, R, P, homol, tablas_ceco, nombre_ceco, R_u=None, P_u=None):
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb['Ppto27']

    # ---------------------------------------------------------- leyenda de vehículos (H3:I8)
    comps = sorted(leyenda.VEHICULOS_PAIS.get(pais, [])) or sorted(set(R['Compañía']) | set(P['Compañía']))
    for r in range(3, 9):
        ws.cell(r, 8).value = None; ws.cell(r, 9).value = None
    for i, c in enumerate(comps[:6]):
        ws.cell(3 + i, 8).value = c
        ws.cell(3 + i, 9).value = leyenda.NOMBRES.get(c, VEHICULOS.get(pais, {}).get(c))
        copiar_estilo(ws.cell(3, 8), ws.cell(3 + i, 8)); copiar_estilo(ws.cell(3, 9), ws.cell(3 + i, 9))

    # ---------------------------------------------------------- líneas
    L = pd.DataFrame()
    meses = []
    if len(R):
        X = preparar(R, homol)
        X['CeCo_txt'] = X.CeCo.map(ceco_norm)
        falt = X.CeCo_txt.isin(['nan', 'None', '', '0'])
        _nc = {str(n).strip().upper(): c for n, c in nombre_ceco.get(pais, {}).items()}
        X.loc[falt, 'CeCo_txt'] = X.loc[falt, 'Nom_CeCo'].map(lambda n: _nc.get(str(n).strip().upper(), ''))
        X['Ceco - Nombre'] = [tablas_ceco.get(c) or (f'{c} - {n}' if c else str(n)) for c, n in zip(X.CeCo_txt, X.Nom_CeCo)]
        malos = X[~X['Ceco - Nombre'].astype(str).str.match(r'^\d+ - ')]
        if len(malos):
            raise ValueError(f'{unidad} {pais}: {len(malos)} filas sin CeCo numerado: ' + str(malos[['Nom_CeCo', 'CeCo']].drop_duplicates().head(10).values.tolist()))
        X['Nivel 1'] = X.Nivel_1.str.replace(r'^\d+ - ', '', regex=True)
        k = ['Compañía', 'Ceco - Nombre', 'Nom_Cta_Contable', 'Nivel 1', 'Nivel_2', 'Proveedor_H', 'Glosa', 'Moneda']
        X = agrupar_amplios(X, k, CUENTAS_CLIENTE)
        exigir_sin_varios(X.Proveedor_H, X['Ceco - Nombre'], f'{unidad} {pais}')
        L = X.pivot_table(index=k, columns='Mes', values='Monto', aggfunc='sum', fill_value=0, observed=True).reset_index()
        meses = [c for c in L.columns if isinstance(c, (int, np.integer))]
        L = L[L[meses].abs().sum(axis=1) > 1e-6]
        L['_sin'] = L.Proveedor_H.eq('(sin proveedor)')
        L = L.sort_values(['Ceco - Nombre', 'Nivel_2', 'Nom_Cta_Contable', '_sin', 'Proveedor_H', 'Glosa', 'Compañía']).reset_index(drop=True)

    estilos = {c: copy.copy(ws.cell(FIRST, c)._style) for c in range(1, 42)}
    formulas = {c: ws.cell(FIRST, c).value for c in (26, 39, 40, 41)}
    last_old = ws.max_row
    for r in range(FIRST, last_old + 1):
        for c in range(1, 42):
            ws.cell(r, c).value = None
    n = max(len(L), 1)
    for i in range(n):
        r = FIRST + i
        for c in range(1, 42):
            ws.cell(r, c)._style = copy.copy(estilos[c])
        ws.cell(r, 26).value = '=+SUM(Tabla1[[#This Row],[Ene-26]:[Dic-26]])'
        ws.cell(r, 39).value = '=+SUM(Tabla1[[#This Row],[Ene-27]:[Dic-27]])'
        ws.cell(r, 40).value = f'=IFERROR(AM{r}/Z{r}-1,0)'
        ws.cell(r, 41).value = '=Tabla1[[#This Row],[P27]]-Tabla1[[#This Row],[F26]]'
        if i >= len(L):
            continue
        row = L.iloc[i]
        ws.cell(r, 1).value = PAIS_TXT[pais]
        ws.cell(r, 2).value = row['Compañía']
        ws.cell(r, 3).value = row['Ceco - Nombre']
        ws.cell(r, 5).value = row['Nom_Cta_Contable']
        ws.cell(r, 6).value = row['Nivel 1']
        ws.cell(r, 7).value = row['Nivel_2']
        ws.cell(r, 10).value = row['Proveedor_H']
        ws.cell(r, 11).value = row['Glosa']
        ws.cell(r, 13).value = row['Moneda']
        for m in meses:
            v = float(row[m])
            if abs(v) > 1e-9:
                ws.cell(r, MES_COL[int(m)]).value = v
    last = FIRST + n - 1
    if last_old > last:
        ws.delete_rows(last + 1, last_old - last)
    ws.tables['Tabla1'].ref = f'A{HDR}:AO{last}'

    # ---------------------------------------------------------- BD - Ppto26
    bd = wb['BD - Ppto26']
    cols = [bd.cell(1, c).value for c in range(1, 31)]
    est_bd = {c: copy.copy(bd.cell(2, c)._style) for c in range(1, 31)}
    bd.delete_rows(2, bd.max_row)
    Pw = P[cols].copy()
    for c in ['CeCo', 'CeCo_Destino', 'Cta_Contable', 'Mes']:
        Pw[c] = Pw[c].map(lambda v: None if pd.isna(v) else re.sub(r'\.0$', '', str(v)))
    Pw['Fecha'] = pd.to_datetime(Pw['Fecha'], errors='coerce')
    Pw = Pw.sort_values(['Fecha', 'Compañía', 'CeCo', 'Nivel_2'], kind='stable')
    filas = []
    for i, rec in enumerate(Pw.itertuples(index=False), start=2):
        fila = []
        for c, v in enumerate(rec, start=1):
            if v is pd.NaT or (isinstance(v, float) and np.isnan(v)): v = None
            if isinstance(v, pd.Timestamp): v = v.to_pydatetime()
            if isinstance(v, str) and v in ('nan', 'None'): v = None
            cell = bd.cell(i, c); cell.value = v; cell._style = copy.copy(est_bd[c])
            fila.append(v)
        filas.append(tuple(fila))
    nbd = len(filas) + 1 if filas else 2

    # ---------------------------------------------------------- TD - Ppto26 (valores visibles; la caché se regenera después)
    td_ref = escribir_td(wb['TD - Ppto26'], Pw, bool(filas))

    carpeta = OUT + ('UdN/' if tipo == 'udn' else 'Staff/') + unidad + '/'
    os.makedirs(carpeta, exist_ok=True)
    nombre = f'Ppto27 - {unidad} - {PAIS_TXT[pais]} - OFICIAL.xlsx'
    wb.save(carpeta + nombre)

    # caché, registros y diseño de la tabla dinámica
    with zipfile.ZipFile(carpeta + nombre) as z:
        cache_id = etree.fromstring(z.read('xl/pivotTables/pivotTable1.xml')).get('cacheId')
    cache_xml, rec_xml, items, idx = pivot_rebuild.build_xml(filas, cols, f'A1:AD{nbd}')
    rows_for_layout = filas if filas else [tuple([None] * 30)]
    ptx = etree.fromstring(pivot_rebuild.pivot_xml(items, rows_for_layout, idx, td_ref))
    ptx.set('cacheId', cache_id)
    partes = pivot_rebuild.partes_logo(carpeta + nombre)
    partes.update({
        'xl/pivotCache/pivotCacheDefinition1.xml': cache_xml,
        'xl/pivotCache/pivotCacheRecords1.xml': rec_xml,
        'xl/pivotTables/pivotTable1.xml': etree.tostring(ptx, xml_declaration=True, encoding='UTF-8', standalone=True),
    })
    pivot_rebuild.reemplazar_partes(carpeta + nombre, partes)
    return carpeta + nombre, L, meses


def escribir_td(ws, Pw, hay_datos):
    """Escribe los valores visibles de la TD (misma disposición que la plantilla) y devuelve su rango."""
    st = {k: copy.copy(ws[k]._style) for k in ['B7', 'C7', 'C8', 'O8', 'B9', 'C9', 'O9', 'B10', 'C10', 'O10', 'B11', 'C11', 'O11', 'B17', 'C17', 'O17']}
    for row in ws.iter_rows(min_row=7, max_row=ws.max_row):
        for c in row:
            c.value = None; c._style = copy.copy(ws['A1']._style)
    fechas = sorted(Pw['Fecha'].dropna().unique()) if hay_datos else []
    nd = max(len(fechas), 1)
    last_col = 3 + nd                  # columna del total
    def put(r, c, v, key):
        cell = ws.cell(r, c); cell.value = v; cell._style = copy.copy(st[key])
    put(7, 2, 'Suma de Monto', 'B7'); put(7, 3, 'Etiquetas de columna', 'C7')
    put(8, 3, 'Ppto 2026' if hay_datos else '(en blanco)', 'C8'); put(8, last_col, 'Total Ppto 2026' if hay_datos else 'Total (en blanco)', 'O8')
    for c in range(4, last_col): ws.cell(8, c)._style = copy.copy(st['C8'])
    put(9, 2, 'Etiquetas de fila', 'B9')
    for j, f in enumerate(fechas):
        put(9, 3 + j, pd.Timestamp(f).to_pydatetime(), 'C9')
    if not fechas: put(9, 3, '(en blanco)', 'C9')
    put(9, last_col, None, 'O9')
    r = 10
    if hay_datos:
        piv = Pw.pivot_table(index=['Nivel_1', 'Nivel_2'], columns='Fecha', values='Monto', aggfunc='sum')
        for n1, g in piv.groupby(level=0, sort=True):
            put(r, 2, n1, 'B10')
            tot = g.sum()
            for j, f in enumerate(fechas):
                put(r, 3 + j, float(tot.get(f, 0)) if f in tot.index else None, 'C10')
            put(r, last_col, float(tot.sum()), 'O10'); r += 1
            for (a, n2), vals in g.sort_index(level=1, key=lambda s: s.astype(str)).iterrows():
                put(r, 2, n2, 'B11')
                for j, f in enumerate(fechas):
                    v = vals.get(f)
                    put(r, 3 + j, None if pd.isna(v) else float(v), 'C11')
                put(r, last_col, float(vals.sum()), 'O11'); r += 1
        tot = Pw.groupby('Fecha').Monto.sum()
        put(r, 2, 'Total general', 'B17')
        for j, f in enumerate(fechas):
            put(r, 3 + j, float(tot.get(f, 0)), 'C17')
        put(r, last_col, float(tot.sum()), 'O17')
    else:
        put(r, 2, '(en blanco)', 'B10'); r += 1
        put(r, 2, '(en blanco)', 'B11'); r += 1
        put(r, 2, 'Total general', 'B17')
        ws.cell(r + 2, 2).value = 'Sin presupuesto 2026 registrado en la TD Regional para esta unidad y país.'
    return f'B7:{get_column_letter(last_col)}{r}'
